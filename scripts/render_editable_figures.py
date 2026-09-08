"""Render the two PhaseSet draw.io pages as native editable PowerPoint shapes.

This is a deliberately small renderer for the closed shapes used by this
project, not a generic draw.io importer. Publication PDF/SVG exports remain
the responsibility of the native draw.io CLI. No slide-wide raster is used.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
from html.parser import HTMLParser
import io
import json
import math
import platform
import posixpath
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import zipfile
import zlib

import lxml.etree
import pptx
import pptx.api
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt


PYTHON_PPTX_VERSION = "1.0.2"
DEFAULT_TEMPLATE_BYTES = 34_030
DEFAULT_TEMPLATE_SHA256 = "e10cc9e120961f6bd4074a373c9c80d2a06c497157e8f4972977b7bea83a8f34"
REGISTERED_SOURCES = {
    "phaseset_motivation.drawio": (
        "a0a25fbc8c55575d50c6ffa8d6add01ce780cc8842328bb3f6bba20c6509d18a"
    ),
    "phaseset_architecture.drawio": (
        "1a61a503c935a72cbfbe1ce175ab54bf975eb3e0a2164018aeda80448f472a1a"
    ),
}
REGISTERED_RUNTIME: dict[str, object] = {
    "libxml2": [2, 11, 9],
    "libxslt": [1, 1, 45],
    "lxml": "6.1.1.0",
    "machine": "AMD64",
    "python": "3.12.14",
    "python_implementation": "CPython",
    "python_pptx": PYTHON_PPTX_VERSION,
    "system": "Windows",
    "zlib_build": "1.3.2",
    "zlib_runtime": "1.3.2",
}
ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)
ZIP_CREATE_SYSTEM = 0
ZIP_EXTERNAL_ATTR = 0o600 << 16

REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
APP_NS = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
DML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

THUMBNAIL_REL = "http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail"
PRINTER_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/printerSettings"
REMOVED_TEMPLATE_PARTS = frozenset(
    {"docProps/thumbnail.jpeg", "ppt/printerSettings/printerSettings1.bin"}
)
RETAINED_PACKAGE_PARTS = frozenset(
    {
        "[Content_Types].xml",
        "_rels/.rels",
        "docProps/app.xml",
        "docProps/core.xml",
        "ppt/_rels/presentation.xml.rels",
        "ppt/presProps.xml",
        "ppt/presentation.xml",
        "ppt/slideMasters/_rels/slideMaster1.xml.rels",
        "ppt/slideMasters/slideMaster1.xml",
        "ppt/slides/_rels/slide1.xml.rels",
        "ppt/slides/slide1.xml",
        "ppt/tableStyles.xml",
        "ppt/theme/theme1.xml",
        "ppt/viewProps.xml",
    }
    | {
        f"ppt/slideLayouts/slideLayout{index}.xml"
        for index in range(1, 12)
    }
    | {
        f"ppt/slideLayouts/_rels/slideLayout{index}.xml.rels"
        for index in range(1, 12)
    }
)
TEMPLATE_OUTPUT_PARTS = RETAINED_PACKAGE_PARTS | REMOVED_TEMPLATE_PARTS

ALLOWED_RELATIONSHIP_TYPES = frozenset(
    {
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/presProps",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/tableStyles",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/viewProps",
        "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties",
    }
)
EXPECTED_RELATIONSHIPS = frozenset(
    {
        (
            "",
            "rId1",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument",
            "ppt/presentation.xml",
        ),
        (
            "",
            "rId3",
            "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties",
            "docProps/core.xml",
        ),
        (
            "",
            "rId4",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties",
            "docProps/app.xml",
        ),
        (
            "ppt/presentation.xml",
            "rId1",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster",
            "slideMasters/slideMaster1.xml",
        ),
        (
            "ppt/presentation.xml",
            "rId3",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/presProps",
            "presProps.xml",
        ),
        (
            "ppt/presentation.xml",
            "rId4",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/viewProps",
            "viewProps.xml",
        ),
        (
            "ppt/presentation.xml",
            "rId5",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme",
            "theme/theme1.xml",
        ),
        (
            "ppt/presentation.xml",
            "rId6",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/tableStyles",
            "tableStyles.xml",
        ),
        (
            "ppt/presentation.xml",
            "rId7",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide",
            "slides/slide1.xml",
        ),
        (
            "ppt/slides/slide1.xml",
            "rId1",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout",
            "../slideLayouts/slideLayout7.xml",
        ),
        (
            "ppt/slideMasters/slideMaster1.xml",
            "rId12",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme",
            "../theme/theme1.xml",
        ),
    }
    | {
        (
            f"ppt/slideLayouts/slideLayout{index}.xml",
            "rId1",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster",
            "../slideMasters/slideMaster1.xml",
        )
        for index in range(1, 12)
    }
    | {
        (
            "ppt/slideMasters/slideMaster1.xml",
            f"rId{index}",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout",
            f"../slideLayouts/slideLayout{index}.xml",
        )
        for index in range(1, 12)
    }
)
DEFAULT_CONTENT_TYPES = {
    "rels": "application/vnd.openxmlformats-package.relationships+xml",
    "xml": "application/xml",
}
OVERRIDE_CONTENT_TYPES = {
    "/docProps/app.xml": "application/vnd.openxmlformats-officedocument.extended-properties+xml",
    "/docProps/core.xml": "application/vnd.openxmlformats-package.core-properties+xml",
    "/ppt/presProps.xml": "application/vnd.openxmlformats-officedocument.presentationml.presProps+xml",
    "/ppt/presentation.xml": "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
    "/ppt/slideMasters/slideMaster1.xml": "application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml",
    "/ppt/slides/slide1.xml": "application/vnd.openxmlformats-officedocument.presentationml.slide+xml",
    "/ppt/tableStyles.xml": "application/vnd.openxmlformats-officedocument.presentationml.tableStyles+xml",
    "/ppt/theme/theme1.xml": "application/vnd.openxmlformats-officedocument.theme+xml",
    "/ppt/viewProps.xml": "application/vnd.openxmlformats-officedocument.presentationml.viewProps+xml",
    **{
        f"/ppt/slideLayouts/slideLayout{index}.xml": (
            "application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"
        )
        for index in range(1, 12)
    },
}
PRIVATE_PATH_RE = re.compile(
    rb"(?:(?<![A-Za-z0-9])[A-Za-z]:[\\/]|/(?:home|media|Users)/|file:|javascript:)",
    re.IGNORECASE,
)


def sha256_bytes(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def xml_bytes(root: ET.Element, default_namespace: str) -> bytes:
    ET.register_namespace("", default_namespace)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def runtime_receipt() -> tuple[dict[str, object], Path]:
    renderer = Path(__file__).resolve(strict=True)
    template = Path(pptx.api.__file__).resolve(strict=True).parent / "templates" / "default.pptx"
    template = template.resolve(strict=True)
    template_raw = template.read_bytes()
    if len(template_raw) != DEFAULT_TEMPLATE_BYTES or sha256_bytes(template_raw) != DEFAULT_TEMPLATE_SHA256:
        raise RuntimeError("python-pptx default template identity mismatch")
    runtime: dict[str, object] = {
        "libxml2": list(lxml.etree.LIBXML_VERSION),
        "libxslt": list(lxml.etree.LIBXSLT_VERSION),
        "lxml": ".".join(str(value) for value in lxml.etree.LXML_VERSION),
        "machine": platform.machine(),
        "python": platform.python_version(),
        "python_implementation": platform.python_implementation(),
        "python_pptx": pptx.__version__,
        "system": platform.system(),
        "zlib_build": zlib.ZLIB_VERSION,
        "zlib_runtime": zlib.ZLIB_RUNTIME_VERSION,
    }
    if runtime != REGISTERED_RUNTIME:
        raise RuntimeError("registered canonical-build runtime identity mismatch")
    receipt = {
        "renderer": {
            "bytes": renderer.stat().st_size,
            "name": renderer.name,
            "sha256": sha256_bytes(renderer.read_bytes()),
        },
        "runtime": runtime,
        "template": {
            "bytes": len(template_raw),
            "name": "python-pptx-default.pptx",
            "sha256": sha256_bytes(template_raw),
        },
    }
    return receipt, template


class LabelParser(HTMLParser):
    def __init__(self, label: str) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[tuple[str, bool]] = []
        self.subscript = False
        self.feed(label)

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        if tag == "sub":
            self.subscript = True
        elif tag == "br":
            self.parts.append(("\n", False))
        else:
            raise ValueError(f"unsupported label tag: {tag}")

    def handle_endtag(self, tag: str) -> None:
        if tag != "sub":
            raise ValueError(f"unsupported closing label tag: {tag}")
        self.subscript = False

    def handle_data(self, data: str) -> None:
        self.parts.append((data, self.subscript))


def style(cell: ET.Element) -> dict[str, str]:
    result = {}
    for field in cell.get("style", "").split(";"):
        if field:
            key, _, value = field.partition("=")
            result[key] = value
    return result


def read_zip_parts(raw: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)) or set(names) != TEMPLATE_OUTPUT_PARTS:
            raise ValueError("python-pptx template part census mismatch")
        if archive.comment:
            raise ValueError("source package comment is not permitted")
        return {name: archive.read(name) for name in names}


def remove_relationship(parts: dict[str, bytes], rels_name: str, relationship_type: str) -> None:
    root = ET.fromstring(parts[rels_name])
    matches = [
        item
        for item in list(root)
        if item.tag == f"{{{REL_NS}}}Relationship"
        and item.get("Type") == relationship_type
    ]
    if len(matches) != 1:
        raise ValueError("template relationship census mismatch")
    root.remove(matches[0])
    parts[rels_name] = xml_bytes(root, REL_NS)


def sanitize_package(raw: bytes) -> dict[str, bytes]:
    parts = read_zip_parts(raw)
    remove_relationship(parts, "_rels/.rels", THUMBNAIL_REL)
    remove_relationship(parts, "ppt/_rels/presentation.xml.rels", PRINTER_REL)
    for name in REMOVED_TEMPLATE_PARTS:
        parts.pop(name)

    content_root = ET.fromstring(parts["[Content_Types].xml"])
    removed_defaults: dict[str, str] = {}
    for item in list(content_root):
        if item.tag != f"{{{CONTENT_TYPES_NS}}}Default":
            continue
        extension = item.get("Extension", "")
        if extension in {"bin", "jpeg"}:
            removed_defaults[extension] = item.get("ContentType", "")
            content_root.remove(item)
    if removed_defaults != {
        "bin": "application/vnd.openxmlformats-officedocument.presentationml.printerSettings",
        "jpeg": "image/jpeg",
    }:
        raise ValueError("template content-type census mismatch")
    parts["[Content_Types].xml"] = xml_bytes(content_root, CONTENT_TYPES_NS)

    app_root = ET.fromstring(parts["docProps/app.xml"])
    stale_fields = {"TotalTime", "Words", "Paragraphs", "HeadingPairs", "TitlesOfParts"}
    for item in list(app_root):
        if item.tag.rsplit("}", 1)[-1] in stale_fields:
            app_root.remove(item)
    truthful = {
        "Application": "PhaseSet native figure renderer",
        "AppVersion": "2.0000",
        "PresentationFormat": "Custom",
        "Slides": "1",
        "Notes": "0",
        "HiddenSlides": "0",
        "MMClips": "0",
        "HyperlinkBase": "",
        "HyperlinksChanged": "false",
        "LinksUpToDate": "false",
        "SharedDoc": "false",
    }
    for name, value in truthful.items():
        item = app_root.find(f"{{{APP_NS}}}{name}")
        if item is None:
            item = ET.SubElement(app_root, f"{{{APP_NS}}}{name}")
        item.text = value
    parts["docProps/app.xml"] = xml_bytes(app_root, APP_NS)
    if set(parts) != RETAINED_PACKAGE_PARTS:
        raise ValueError("retained package part census mismatch")
    return parts


def normalized_zip(parts: dict[str, bytes]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as target:
        for name in sorted(parts):
            info = zipfile.ZipInfo(name, ZIP_TIMESTAMP)
            info.compress_type = zipfile.ZIP_STORED
            info.create_system = ZIP_CREATE_SYSTEM
            info.create_version = 20
            info.extract_version = 20
            info.flag_bits = 0
            info.external_attr = ZIP_EXTERNAL_ATTR
            info.internal_attr = 0
            info.extra = b""
            info.comment = b""
            target.writestr(info, parts[name], compress_type=zipfile.ZIP_STORED)
    return output.getvalue()


def relationship_owner(rels_name: str) -> str:
    if rels_name == "_rels/.rels":
        return ""
    prefix, local_name = rels_name.rsplit("/_rels/", 1)
    if not local_name.endswith(".rels"):
        raise ValueError("malformed relationship part name")
    return f"{prefix}/{local_name[:-5]}"


def validate_content_types(parts: dict[str, bytes]) -> None:
    root = ET.fromstring(parts["[Content_Types].xml"])
    if root.tag != f"{{{CONTENT_TYPES_NS}}}Types" or any(
        item.tag
        not in {
            f"{{{CONTENT_TYPES_NS}}}Default",
            f"{{{CONTENT_TYPES_NS}}}Override",
        }
        for item in root
    ):
        raise ValueError("retained content-type schema mismatch")
    default_rows = root.findall(f"{{{CONTENT_TYPES_NS}}}Default")
    override_rows = root.findall(f"{{{CONTENT_TYPES_NS}}}Override")
    defaults = {
        item.get("Extension", ""): item.get("ContentType", "")
        for item in default_rows
    }
    overrides = {
        item.get("PartName", ""): item.get("ContentType", "")
        for item in override_rows
    }
    if (
        len(defaults) != len(default_rows)
        or len(overrides) != len(override_rows)
        or defaults != DEFAULT_CONTENT_TYPES
        or overrides != OVERRIDE_CONTENT_TYPES
    ):
        raise ValueError("retained content-type allowlist mismatch")


def validate_relationships(parts: dict[str, bytes]) -> None:
    relationship_rows: set[tuple[str, str, str, str]] = set()
    for name in sorted(item for item in parts if item.endswith(".rels")):
        root = ET.fromstring(parts[name])
        if root.tag != f"{{{REL_NS}}}Relationships" or any(
            item.tag != f"{{{REL_NS}}}Relationship" for item in root
        ):
            raise ValueError("retained relationship schema mismatch")
        identifiers: set[str] = set()
        for relationship in root.findall(f"{{{REL_NS}}}Relationship"):
            identifier = relationship.get("Id", "")
            relation_type = relationship.get("Type", "")
            target = relationship.get("Target", "")
            if (
                not re.fullmatch(r"rId[1-9][0-9]*", identifier)
                or identifier in identifiers
                or relation_type not in ALLOWED_RELATIONSHIP_TYPES
                or relationship.get("TargetMode") is not None
                or not target
                or "\\" in target
                or ":" in target
                or target.startswith("/")
            ):
                raise ValueError("retained relationship allowlist mismatch")
            identifiers.add(identifier)
            owner = relationship_owner(name)
            relationship_rows.add((owner, identifier, relation_type, target))
            resolved = posixpath.normpath(posixpath.join(posixpath.dirname(owner), target))
            if resolved.startswith("../") or resolved not in parts:
                raise ValueError("retained relationship target escapes package")
    if relationship_rows != EXPECTED_RELATIONSHIPS:
        raise ValueError("retained relationship tuple census mismatch")


def validate_app_properties(parts: dict[str, bytes]) -> None:
    root = ET.fromstring(parts["docProps/app.xml"])
    expected = {
        "Application": "PhaseSet native figure renderer",
        "AppVersion": "2.0000",
        "PresentationFormat": "Custom",
        "Slides": "1",
        "Notes": "0",
        "HiddenSlides": "0",
        "MMClips": "0",
    }
    for name, value in expected.items():
        item = root.find(f"{{{APP_NS}}}{name}")
        if item is None or (item.text or "") != value:
            raise ValueError("extended property truthfulness mismatch")
    for name in ("TotalTime", "Words", "Paragraphs", "HeadingPairs", "TitlesOfParts"):
        if root.find(f"{{{APP_NS}}}{name}") is not None:
            raise ValueError("stale extended property retained")


def validate_hardened_package(
    raw: bytes,
    expected_shape_names: set[str],
    expected_connector_names: set[str],
) -> tuple[int, int]:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        names = archive.namelist()
        if names != sorted(RETAINED_PACKAGE_PARTS) or len(names) != len(set(names)):
            raise ValueError("normalized package part order or census mismatch")
        if archive.comment:
            raise ValueError("normalized package comment is not permitted")
        for info in archive.infolist():
            if (
                info.date_time != ZIP_TIMESTAMP
                or info.compress_type != zipfile.ZIP_STORED
                or info.create_system != ZIP_CREATE_SYSTEM
                or info.create_version != 20
                or info.extract_version != 20
                or info.flag_bits != 0
                or info.external_attr != ZIP_EXTERNAL_ATTR
                or info.internal_attr != 0
                or info.extra
                or info.comment
            ):
                raise ValueError("normalized ZIP metadata mismatch")
        parts = {name: archive.read(name) for name in names}
    for name, part in parts.items():
        upper = part.upper()
        if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper or PRIVATE_PATH_RE.search(part):
            raise ValueError(f"unsafe retained package text in {name}")
    validate_content_types(parts)
    validate_relationships(parts)
    validate_app_properties(parts)

    slide_xml = ET.fromstring(parts["ppt/slides/slide1.xml"])
    ns = {"p": PML_NS, "a": DML_NS}
    shape_names = [
        item.get("name", "")
        for item in slide_xml.findall(".//p:sp/p:nvSpPr/p:cNvPr", ns)
    ]
    connector_names = [
        item.get("name", "")
        for item in slide_xml.findall(".//p:cxnSp/p:nvCxnSpPr/p:cNvPr", ns)
    ]
    if (
        len(shape_names) != len(set(shape_names))
        or set(shape_names) != expected_shape_names
        or len(connector_names) != len(set(connector_names))
        or set(connector_names) != expected_connector_names
    ):
        raise ValueError("native object identity census is incomplete")
    forbidden_tags = {
        "pic", "graphicFrame", "oleObj", "control", "contentPart", "hlinkClick",
        "hlinkHover", "video", "audio",
    }
    if any(item.tag.rsplit("}", 1)[-1] in forbidden_tags for item in slide_xml.iter()):
        raise ValueError("forbidden slide object retained")
    return len(shape_names), len(connector_names)


def render(source: Path, destination: Path, template: Path) -> dict[str, object]:
    raw = source.read_bytes()
    source_sha256 = sha256_bytes(raw)
    if REGISTERED_SOURCES.get(source.name) != source_sha256:
        raise ValueError("registered figure source identity mismatch")
    if (
        len(raw) > 1024 * 1024
        or b"<!DOCTYPE" in raw.upper()
        or b"<!ENTITY" in raw.upper()
        or PRIVATE_PATH_RE.search(raw)
        or re.search(rb"https?://", raw, re.IGNORECASE)
    ):
        raise ValueError("unsupported source document")
    tree = ET.fromstring(raw)
    models = tree.findall("./diagram/mxGraphModel")
    if len(models) != 1:
        raise ValueError("exactly one uncompressed page is required")
    model = models[0]
    cell_rows = model.findall("./root/mxCell")
    cell_ids = [cell.attrib["id"] for cell in cell_rows]
    if len(cell_ids) != len(set(cell_ids)):
        raise ValueError("duplicate draw.io cell id")
    cells = dict(zip(cell_ids, cell_rows, strict=True))
    width = float(model.attrib["pageWidth"])
    height = float(model.attrib["pageHeight"])
    prs = Presentation(str(template))
    prs.slide_width = Inches(16)
    scale = int(prs.slide_width) / width
    prs.slide_height = Emu(round(height * scale))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    properties = prs.core_properties
    properties.title = source.stem
    properties.author = "PhaseSet contributors"
    properties.last_modified_by = "PhaseSet figure renderer"
    properties.subject = "Data-free methods figure; native editable objects"
    properties.comments = "Source SHA256: " + source_sha256
    properties.created = properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    boxes: dict[str, tuple[float, float, float, float]] = {}

    def box(identifier: str) -> tuple[float, float, float, float]:
        if identifier in boxes:
            return boxes[identifier]
        cell = cells[identifier]
        geometry = cell.find("mxGeometry")
        if geometry is None:
            return (0, 0, 0, 0)
        parent = cell.get("parent", "1")
        px, py, _, _ = box(parent) if parent in cells and parent != identifier else (0, 0, 0, 0)
        result = (
            px + float(geometry.get("x", "0")),
            py + float(geometry.get("y", "0")),
            float(geometry.get("width", "0")),
            float(geometry.get("height", "0")),
        )
        boxes[identifier] = result
        return result

    def emu(value: float) -> Emu:
        return Emu(round(value * scale))

    def font_points(value: float) -> Pt:
        return Pt(value * scale / 12700)

    def format_line(shape, settings: dict[str, str]) -> None:
        shape.line.color.rgb = RGBColor.from_string(settings.get("strokeColor", "#000000").lstrip("#"))
        shape.line.width = font_points(float(settings.get("strokeWidth", "1")))
        if settings.get("dashed") == "1":
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    def label(shape, text: str, settings: dict[str, str]) -> None:
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = emu(3)
        frame.margin_top = frame.margin_bottom = emu(1)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.space_before = paragraph.space_after = Pt(0)
        for text_part, subscript in LabelParser(text).parts:
            for index, segment in enumerate(text_part.split("\n")):
                if index:
                    paragraph = frame.add_paragraph()
                    paragraph.alignment = PP_ALIGN.CENTER
                    paragraph.space_before = paragraph.space_after = Pt(0)
                if segment:
                    run = paragraph.add_run()
                    run.text = segment
                    run.font.name = "Arial"
                    size = float(settings.get("fontSize", "14"))
                    run.font.size = font_points(size * (0.75 if subscript else 1))
                    run.font.bold = bool(int(settings.get("fontStyle", "0")) & 1)
                    run.font.color.rgb = RGBColor.from_string(settings.get("fontColor", "#000000").lstrip("#"))
                    if subscript:
                        run._r.get_or_add_rPr().set("baseline", "-25000")

    vertices = [cell for cell in cells.values() if cell.get("vertex") == "1"]
    edges = [cell for cell in cells.values() if cell.get("edge") == "1"]
    containers = [cell for cell in vertices if "swimlane" in style(cell)]
    expected_shape_names = {
        *("drawio-node-" + cell.attrib["id"] for cell in vertices),
        *("drawio-header-" + cell.attrib["id"] for cell in containers),
        *(
            "drawio-edge-label-" + cell.attrib["id"]
            for cell in edges
            if cell.get("value")
        ),
    }
    expected_connector_names: set[str] = set()
    edge_segment_counts = {cell.attrib["id"]: 0 for cell in edges}

    def vertex(cell: ET.Element, *, container: bool = False) -> None:
        identifier = cell.attrib["id"]
        settings = style(cell)
        x, y, w, h = box(identifier)
        if "text" in settings:
            shape = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
        else:
            kind = (
                MSO_AUTO_SHAPE_TYPE.OVAL if "ellipse" in settings else
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if settings.get("rounded") == "1" else
                MSO_AUTO_SHAPE_TYPE.RECTANGLE
            )
            shape = slide.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(settings.get("fillColor", "#FFFFFF").lstrip("#"))
            format_line(shape, settings)
            if kind == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE:
                shape.adjustments[0] = 0.08
        shape.name = "drawio-node-" + identifier
        text = cell.get("value", "")
        if container:
            header = float(settings.get("startSize", "36"))
            title = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(header))
            title.name = "drawio-header-" + identifier
            label(title, text, settings)
            separator = slide.shapes.add_connector(MSO_CONNECTOR_TYPE.STRAIGHT, emu(x), emu(y + header), emu(x + w), emu(y + header))
            separator.name = "drawio-container-separator-" + identifier
            expected_connector_names.add(separator.name)
            format_line(separator, settings)
        else:
            label(shape, text, settings)

    for cell in containers:
        vertex(cell, container=True)

    def endpoint(identifier: str, other: str, edge_style: dict[str, str], prefix: str) -> tuple[float, float]:
        x, y, w, h = box(identifier)
        if prefix + "X" in edge_style:
            return x + float(edge_style[prefix + "X"]) * w, y + float(edge_style[prefix + "Y"]) * h
        ox, oy, ow, oh = box(other)
        dx, dy = ox + ow / 2 - x - w / 2, oy + oh / 2 - y - h / 2
        if "ellipse" in style(cells[identifier]):
            factor = 1 / math.sqrt((2 * dx / w) ** 2 + (2 * dy / h) ** 2)
        else:
            factor = 1 / max(abs(2 * dx / w), abs(2 * dy / h))
        return x + w / 2 + dx * factor, y + h / 2 + dy * factor

    for cell in edges:
        settings = style(cell)
        source_id, target_id = cell.attrib["source"], cell.attrib["target"]
        start = endpoint(source_id, target_id, settings, "exit")
        end = endpoint(target_id, source_id, settings, "entry")
        geometry = cell.find("mxGeometry")
        px, py, _, _ = box(cell.get("parent", "1"))
        points = [] if geometry is None else [
            (px + float(point.attrib["x"]), py + float(point.attrib["y"]))
            for point in geometry.findall("./Array/mxPoint")
        ]
        if not points and settings.get("edgeStyle") == "orthogonalEdgeStyle" and start[0] != end[0] and start[1] != end[1]:
            middle = (start[0] + end[0]) / 2
            points = [(middle, start[1]), (middle, end[1])]
        route = [start, *points, end]
        for index, (a, b) in enumerate(zip(route, route[1:])):
            if a == b:
                continue
            connector = slide.shapes.add_connector(MSO_CONNECTOR_TYPE.STRAIGHT, emu(a[0]), emu(a[1]), emu(b[0]), emu(b[1]))
            connector.name = f"drawio-edge-{cell.attrib['id']}-segment-{index}"
            expected_connector_names.add(connector.name)
            edge_segment_counts[cell.attrib["id"]] += 1
            format_line(connector, settings)
            if index == len(route) - 2 and settings.get("endArrow", "none") != "none":
                tail = OxmlElement("a:tailEnd")
                tail.set("type", "triangle")
                connector._element.spPr.get_or_add_ln().append(tail)
        if cell.get("value"):
            a, b = route[len(route) // 2 - 1:len(route) // 2 + 1]
            tx, ty = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2
            tag = slide.shapes.add_textbox(emu(tx - 115), emu(ty - 12), emu(230), emu(24))
            tag.name = "drawio-edge-label-" + cell.attrib["id"]
            tag.fill.solid()
            tag.fill.fore_color.rgb = RGBColor(255, 255, 255)
            label(tag, cell.attrib["value"], settings)
    for cell in vertices:
        if cell not in containers:
            vertex(cell)
    if any(count < 1 for count in edge_segment_counts.values()):
        raise ValueError("an editable edge produced no native connector segment")
    output = io.BytesIO()
    prs.save(output)
    normalized = normalized_zip(sanitize_package(output.getvalue()))
    shapes, connectors = validate_hardened_package(
        normalized,
        expected_shape_names,
        expected_connector_names,
    )
    destination.write_bytes(normalized)
    return {
        "source": source.name, "source_sha256": source_sha256,
        "pptx": destination.name, "pptx_sha256": hashlib.sha256(normalized).hexdigest(),
        "native_shapes": shapes, "native_connector_segments": connectors,
        "native_node_shapes": len(vertices),
        "native_container_headers": len(containers),
        "native_edge_labels": sum(bool(cell.get("value")) for cell in edges),
        "source_edges": len(edges),
        "package_parts": len(RETAINED_PACKAGE_PARTS),
        "external_relationships": 0,
        "raster_media": 0,
        "empirical_result": False,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--figure-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    destinations = tuple(
        args.output_dir / name
        for name in (
            "phaseset_motivation.pptx",
            "phaseset_architecture.pptx",
            "EDITABLE_FIGURES.json",
        )
    )
    if any(path.exists() for path in destinations):
        raise FileExistsError("refusing to overwrite an existing figure receipt")
    receipt, template = runtime_receipt()
    rows = []
    for stem in ("phaseset_motivation", "phaseset_architecture"):
        rows.append(
            render(
                args.figure_dir / (stem + ".drawio"),
                args.output_dir / (stem + ".pptx"),
                template,
            )
        )
    result = {
        "schema": "phaseset-editable-figure-census-v2",
        "figures": rows,
        **receipt,
        "package_policy": {
            "compression": "ZIP_STORED",
            "create_system": ZIP_CREATE_SYSTEM,
            "external_attr": ZIP_EXTERNAL_ATTR,
            "retained_part_count": len(RETAINED_PACKAGE_PARTS),
            "timestamp": list(ZIP_TIMESTAMP),
        },
    }
    (args.output_dir / "EDITABLE_FIGURES.json").write_text(json.dumps(result, sort_keys=True, indent=2) + "\n", encoding="ascii")
    print(json.dumps(result, sort_keys=True))


if __name__ == "__main__":
    main()
